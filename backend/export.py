"""Сборка официального протокола совещания в DOCX.

На выходе документ, который можно подписать и разослать:
шапка, участники, ход совещания, таблица поручений.
"""

from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm


def _style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def build_protocol(
    path: str,
    title: str,
    meeting_date: date,
    participants: list[str],
    transcript: list[dict],
    tasks: list[dict],
    summary: str = "",
    organization: str = "",
) -> str:
    doc = Document()
    _style(doc)
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)

    if organization:
        p = doc.add_paragraph(organization)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].bold = True

    h = doc.add_paragraph("ПРОТОКОЛ СОВЕЩАНИЯ")
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.runs[0].bold = True
    h.runs[0].font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"от {meeting_date:%d.%m.%Y}")

    doc.add_paragraph().add_run(f"Тема: {title}").bold = True

    if participants:
        doc.add_paragraph("Участники:").runs[0].bold = True
        for name in participants:
            doc.add_paragraph(name, style="List Bullet")

    if summary:
        doc.add_paragraph("Краткое содержание:").runs[0].bold = True
        doc.add_paragraph(summary)

    doc.add_paragraph("Поручения:").runs[0].bold = True
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, name in enumerate(["№", "Поручение", "Ответственный", "Срок", "Статус"]):
        cell = table.rows[0].cells[i]
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True

    for n, t in enumerate(tasks, 1):
        row = table.add_row().cells
        row[0].text = str(n)
        row[1].text = t.get("text", "")
        row[2].text = t.get("assignee", "")
        deadline = t.get("deadline")
        if deadline:
            d = datetime.fromisoformat(deadline).date()
            row[3].text = f"{d:%d.%m.%Y}"
            if t.get("deadline_raw"):
                row[3].paragraphs[0].add_run(f"\n({t['deadline_raw']})").italic = True
        else:
            row[3].text = t.get("deadline_raw", "не назван")
        row[4].text = t.get("status", "в работе")

    if transcript:
        doc.add_page_break()
        doc.add_paragraph("Приложение. Стенограмма совещания").runs[0].bold = True
        for seg in transcript:
            p = doc.add_paragraph()
            p.add_run(f"{seg.get('speaker', 'Участник')}: ").bold = True
            p.add_run(seg.get("text", ""))

    foot = doc.add_paragraph()
    foot.add_run(
        "\nПротокол сформирован автоматически системой расшифровки совещаний. "
        "Требует проверки и утверждения ответственным лицом."
    ).italic = True

    doc.save(path)
    return path


if __name__ == "__main__":
    demo_tasks = [
        {
            "text": "Подготовить претензию поставщику за срыв сроков",
            "assignee": "Ерлан, юрист департамента",
            "deadline_raw": "до конца недели",
            "deadline": "2026-09-25",
            "status": "в работе",
        },
        {
            "text": "Найти альтернативного поставщика сырья для расчёта",
            "assignee": "Ботагоз Нурлановна",
            "deadline_raw": "за две недели",
            "deadline": "2026-10-07",
            "status": "в работе",
        },
    ]
    out = build_protocol(
        "samples/protocol_demo.docx",
        title="Доклады по производственным показателям направлений",
        meeting_date=date(2026, 9, 23),
        participants=["Данияр Серикович", "Ботагоз Нурлановна", "Жандос Талгатович"],
        transcript=[{"speaker": "Данияр Серикович", "text": "Коллеги, начинаем совещание."}],
        tasks=demo_tasks,
        summary="Рассмотрены показатели по направлениям, выданы поручения.",
        organization='АО «Самрук-Қазына Ондеу»',
    )
    print("готово:", out)
